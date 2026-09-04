import csv
import io
import logging
import os
from pathlib import Path

from app.config import Config

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".csv", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".css", ".log", ".env", ".ini", ".toml",
    ".sql", ".sh", ".bat", ".rtf",
}


class FileTool:
    """Sandboxed file operations for user uploads."""

    def __init__(self, db):
        self.db = db
        self.uploads_root = Config.UPLOADS_DIR
        os.makedirs(self.uploads_root, exist_ok=True)

    def _user_dir(self, user_id: int) -> str:
        path = os.path.join(self.uploads_root, str(user_id))
        os.makedirs(path, exist_ok=True)
        return path

    def _safe_path(self, user_id: int, relative_path: str) -> tuple[str | None, str | None]:
        if not relative_path or not relative_path.strip():
            return None, "Путь не указан"
        rel = relative_path.strip().replace("\\", "/").lstrip("/")
        if ".." in rel.split("/"):
            return None, "Доступ запрещён: path traversal"
        user_dir = self._user_dir(user_id)
        full = os.path.abspath(os.path.join(user_dir, rel))
        if not full.startswith(os.path.abspath(user_dir)):
            return None, "Доступ запрещён: выход за пределы папки"
        ext = Path(rel).suffix.lower()
        if ext and ext not in Config.ALLOWED_FILE_EXTENSIONS:
            return None, f"Расширение {ext} не разрешено"
        return full, None

    def _resolve_file(self, user_id: int, file_id: int = None, filename: str = None) -> dict | None:
        if file_id:
            record = self.db.get_user_file(user_id, file_id=file_id)
            if record:
                return record
        if filename:
            record = self.db.get_user_file(user_id, filename=filename)
            if record:
                return record
        return None

    def is_tabular(self, filename: str) -> bool:
        return Path(filename).suffix.lower() in {".csv", ".tsv", ".xlsx", ".xls"}

    def _read_xlsx(self, path: str) -> str:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames[:3]:
            ws = wb[sheet_name]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 200:
                    rows.append("...[строки обрезаны]")
                    break
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                parts.append(f"--- Лист: {sheet_name} ---\n" + "\n".join(rows))
        wb.close()
        return "\n\n".join(parts) or "(пустая таблица)"

    def _read_csv_text(self, path: str) -> tuple[str, dict]:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            sample = f.read()
        delimiter = "\t" if path.endswith(".tsv") else ","
        reader = csv.reader(io.StringIO(sample), delimiter=delimiter)
        rows = list(reader)
        meta = {"rows": len(rows), "cols": len(rows[0]) if rows else 0, "is_table": True}
        if rows:
            meta["headers"] = rows[0]
        return sample, meta

    def _extract_text(self, path: str, mime_type: str = None) -> tuple[str | None, str | None, dict]:
        meta: dict = {"is_table": False}
        ext = Path(path).suffix.lower()
        try:
            if ext == ".pdf":
                try:
                    from pypdf import PdfReader
                except ImportError:
                    return None, "PDF не поддерживается: установите pypdf", meta
                reader = PdfReader(path)
                parts = []
                for page in reader.pages[:30]:
                    parts.append(page.extract_text() or "")
                text = "\n".join(parts).strip()
                return (text or "(PDF без извлекаемого текста)", None, meta)

            if ext == ".xlsx":
                try:
                    text = self._read_xlsx(path)
                    meta["is_table"] = True
                    return text, None, meta
                except ImportError:
                    return None, "Excel не поддерживается: установите openpyxl", meta

            if ext in (".csv", ".tsv"):
                text, meta = self._read_csv_text(path)
                return text, None, meta

            if ext in TEXT_EXTENSIONS or (mime_type and mime_type.startswith("text/")):
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    return f.read(), None, meta

            size = os.path.getsize(path)
            return None, f"Бинарный файл ({ext or mime_type or 'unknown'}, {size} байт). Пришли txt, csv, pdf или xlsx.", meta
        except Exception as e:
            logger.error(f"Text extraction failed for {path}: {e}")
            return None, str(e), meta

    def _truncate(self, text: str) -> str:
        limit = Config.MAX_FILE_READ_CHARS
        if len(text) <= limit:
            return text
        return text[:limit] + f"\n\n...[обрезано, всего {len(text)} символов]"

    def list_files(self, user_id: int) -> dict:
        files = self.db.list_user_files(user_id)
        return {"success": True, "files": files, "count": len(files)}

    def read_file(self, user_id: int, file_id: int = None, filename: str = None) -> dict:
        record = self._resolve_file(user_id, file_id=file_id, filename=filename)
        if not record:
            return {"success": False, "error": "Файл не найден. Используй list_user_files."}
        path = record["local_path"]
        if not os.path.exists(path):
            return {"success": False, "error": "Файл на диске не найден"}
        text, err, meta = self._extract_text(path, record.get("mime_type"))
        if err:
            return {"success": False, "error": err, "file_id": record["id"], "name": record["original_name"]}
        return {
            "success": True,
            "file_id": record["id"],
            "name": record["original_name"],
            "content": self._truncate(text),
            "size": record.get("size"),
            "is_table": meta.get("is_table") or self.is_tabular(record["original_name"]),
            "table_meta": meta,
        }

    def write_file(
        self,
        user_id: int,
        filename: str,
        content: str,
        send_to_chat: bool = False,
    ) -> dict:
        safe_name = os.path.basename(filename.replace("\\", "/"))
        if not safe_name:
            return {"success": False, "error": "Нужно имя файла"}
        full_path, err = self._safe_path(user_id, safe_name)
        if err:
            return {"success": False, "error": err}
        encoded = content.encode("utf-8")
        if len(encoded) > Config.MAX_FILE_SIZE_BYTES:
            return {"success": False, "error": f"Файл слишком большой (лимит {Config.MAX_FILE_SIZE_BYTES // 1024 // 1024} МБ)"}
        os.makedirs(os.path.dirname(full_path) or self._user_dir(user_id), exist_ok=True)
        with open(full_path, "w", encoding="utf-8") as f:
            f.write(content)
        file_id = self.db.add_user_file(
            user_id,
            telegram_file_id=None,
            local_path=full_path,
            original_name=safe_name,
            mime_type="text/plain",
            size=len(encoded),
        )
        result = {
            "success": True,
            "file_id": file_id,
            "name": safe_name,
            "message": f"Файл `{safe_name}` сохранён (id={file_id})",
            "send_to_chat": bool(send_to_chat),
        }
        return result

    def delete_file(self, user_id: int, file_id: int = None, filename: str = None) -> dict:
        record = self._resolve_file(user_id, file_id=file_id, filename=filename)
        if not record:
            return {"success": False, "error": "Файл не найден"}
        path = record["local_path"]
        if os.path.exists(path):
            try:
                os.remove(path)
            except OSError as e:
                return {"success": False, "error": str(e)}
        self.db.delete_user_file(user_id, record["id"])
        return {"success": True, "message": f"Файл `{record['original_name']}` удалён"}

    def register_download(
        self,
        user_id: int,
        local_path: str,
        original_name: str,
        telegram_file_id: str = None,
        mime_type: str = None,
        size: int = None,
    ) -> dict:
        file_id = self.db.add_user_file(
            user_id, telegram_file_id, local_path, original_name, mime_type, size
        )
        preview, _, meta = self._extract_text(local_path, mime_type)
        preview_short = (preview[:500] + "...") if preview and len(preview) > 500 else preview
        return {
            "success": True,
            "file_id": file_id,
            "name": original_name,
            "size": size,
            "preview": preview_short,
        }
